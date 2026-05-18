"""
Tests for the Google Drive integration (DriveAgent + OAuth isolation).

Safety-critical invariants under test:
  - Drive OAuth never touches calendar/connected_accounts state.
  - The edit resolver is deterministic and REFUSES on 0 / >1 matches.
  - propose_modification stages a preview and writes NOTHING.
  - apply_modification only runs on explicit confirmation, and the revision
    guard blocks a write when the file moved since the preview.
  - Routing reaches DriveAgent without disturbing the existing chain.

Firebase is stubbed by the root conftest.py.
"""
from unittest.mock import MagicMock, patch

import pytest

from app.agents.drive_agent import DriveAgent
from app.agents.drive_agent.skill_context import SkillContext
from app.agents.drive_agent._shared.edit_resolver import (
    col_letter,
    resolve_sheet_edit,
    resolve_text_edit,
    validate_edit_spec,
)
from app.db.user_context_store import get_user_context, update_user_context
from app.models.agent_result import AgentResult
from app.models.inbound_message import InboundMessage
from app.models.parsed_message import ParsedMessage
from app.repositories.user_repository import UserRepository
from app.responder.response_formatter import format_response
from app.router.deterministic_router import route
from app.handlers.pending_drive_handler import handle_pending_drive

USER = {"phone_number": "+573001234567", "language": "en",
        "google_drive_refresh_token": "enc"}
GRID = [["cliente", "estado"], ["pepito", "pendiente"], ["ana", "ok"]]
SHEET_MIME = "application/vnd.google-apps.spreadsheet"
DOC_MIME = "application/vnd.google-apps.document"


def _pm(text, **kw):
    return ParsedMessage(raw_message=text, signals=[], **kw)


def _ib(text, phone="+573001234567"):
    return InboundMessage(user_phone_number=phone, text=text,
                          message_id="m1", message_type="text")


# ── OAuth isolation ─────────────────────────────────────────────────────────

def _patch_db(snapshot_dict, exists=True):
    captured = {}
    doc_ref = MagicMock()
    snap = MagicMock()
    snap.exists = exists
    snap.to_dict.return_value = dict(snapshot_dict)
    doc_ref.get.return_value = snap
    doc_ref.set.side_effect = lambda data, **kw: captured.update(data)
    db = MagicMock()
    db.collection.return_value.document.return_value = doc_ref
    return db, captured


def test_drive_oauth_state_does_not_touch_calendar_fields():
    db, captured = _patch_db({})
    with patch("app.repositories.user_repository.db", db):
        UserRepository.set_drive_oauth_state_token("+57", "tok", None, code_verifier="cv")
    assert captured["google_drive_oauth_state_token"] == "tok"
    assert captured["google_drive_oauth_code_verifier"] == "cv"
    # The shared calendar OAuth namespace must be untouched.
    assert "google_oauth_state_token" not in captured
    assert "oauth_pending_provider" not in captured
    assert "connected_accounts" not in captured


def test_save_drive_credentials_isolated():
    db, captured = _patch_db({})
    with patch("app.repositories.user_repository.db", db):
        UserRepository.save_drive_credentials("+57", "ENC")
    assert captured["google_drive_refresh_token"] == "ENC"
    assert captured["google_drive_connected"] is True
    assert "google_calendar_refresh_token" not in captured
    assert "connected_accounts" not in captured


# ── edit resolver determinism ───────────────────────────────────────────────

def test_col_letter():
    assert col_letter(0) == "A"
    assert col_letter(25) == "Z"
    assert col_letter(26) == "AA"


@pytest.mark.parametrize("spec,err", [
    ({"op": "set_cell", "locator_column": "c", "locator_value": "v",
      "target_column": "t", "new_value": "x"}, None),
    ({"op": "bogus"}, "invalid_edit_spec"),
    ({"op": "set_cell", "locator_column": "c"}, "invalid_edit_spec"),
    ({"op": "replace_text", "find": "a", "replace": "b"}, None),
    ({"op": "replace_text", "find": "", "replace": "b"}, "invalid_edit_spec"),
    ({"op": "append_text", "text": "hi"}, None),
    (None, "invalid_edit_spec"),
])
def test_validate_edit_spec(spec, err):
    assert validate_edit_spec(spec) == err


def test_resolve_sheet_edit_single_match_case_insensitive():
    r = resolve_sheet_edit(GRID, {"op": "set_cell", "locator_column": "Cliente",
                                  "locator_value": "PEPITO",
                                  "target_column": "estado", "new_value": "OK"})
    assert r["ok"] and r["a1"] == "B2"
    assert r["old_value"] == "pendiente" and r["new_value"] == "OK"


def test_resolve_sheet_edit_refuses_no_and_multi_match():
    assert resolve_sheet_edit(GRID, {"op": "set_cell", "locator_column": "cliente",
        "locator_value": "zzz", "target_column": "estado",
        "new_value": "x"})["error"] == "edit_no_match"
    g = [["c", "s"], ["dup", "1"], ["dup", "2"]]
    assert resolve_sheet_edit(g, {"op": "set_cell", "locator_column": "c",
        "locator_value": "dup", "target_column": "s",
        "new_value": "x"})["error"] == "edit_multiple_matches"


def test_resolve_sheet_edit_column_not_found():
    r = resolve_sheet_edit(GRID, {"op": "set_cell", "locator_column": "nope",
        "locator_value": "pepito", "target_column": "estado", "new_value": "x"})
    assert r["error"] == "edit_column_not_found"


def test_resolve_text_edit_exact_single_occurrence_only():
    ok = resolve_text_edit("hello world", {"op": "replace_text",
                                           "find": "world", "replace": "there"})
    assert ok["ok"] and ok["new_content"] == "hello there"
    assert resolve_text_edit("a a", {"op": "replace_text", "find": "a",
                                     "replace": "b"})["error"] == "edit_multiple_matches"
    assert resolve_text_edit("x", {"op": "replace_text", "find": "q",
                                   "replace": "b"})["error"] == "edit_no_match"


def test_resolve_text_edit_append():
    r = resolve_text_edit("line1", {"op": "append_text", "text": "line2"})
    assert r["ok"] and r["new_content"] == "line1\nline2"


# ── propose stages a preview, writes nothing ────────────────────────────────

def test_propose_modification_stages_preview_no_write():
    pm = _pm("set estado OK for cliente pepito in Ventas",
             drive_intent="modify", drive_file_ref="Ventas",
             drive_edit={"op": "set_cell", "locator_column": "cliente",
                         "locator_value": "pepito", "target_column": "estado",
                         "new_value": "OK"})
    with patch("app.agents.drive_agent._shared.drive_client.decrypt", return_value="rt"), \
         patch("app.services.google_drive.search_files",
               return_value=[{"id": "S1", "name": "Ventas", "mimeType": SHEET_MIME}]), \
         patch("app.services.google_drive.get_file_meta",
               return_value={"headRevisionId": "rev1"}), \
         patch("app.services.google_drive.read_sheet_values",
               return_value=("Hoja1", GRID)), \
         patch("app.services.google_drive.update_sheet_cell") as wrote:
        res = DriveAgent().execute(pm, USER)

    assert res.success and res.data["type"] == "drive_modify_preview"
    assert res.data["old_value"] == "pendiente" and res.data["new_value"] == "OK"
    wrote.assert_not_called()
    pend = get_user_context(USER["phone_number"])["pending_drive"]
    assert pend["step"] == "awaiting_modify_confirmation"
    assert pend["a1"] == "B2" and pend["expected_revision"] == "rev1"
    msg = format_response(res, USER)
    assert "Ventas" in msg and "pendiente" in msg and "OK" in msg
    update_user_context(USER["phone_number"], "pending_drive", None)


# ── gate: confirmation is mandatory ─────────────────────────────────────────

def _stage(phone="+573001234567"):
    pending = {"step": "awaiting_modify_confirmation", "op": "set_cell",
               "file_name": "Ventas", "spreadsheet_id": "S1",
               "sheet_name": "Hoja1", "a1": "B2", "new_value": "OK",
               "expected_revision": "rev1", "mime_type": SHEET_MIME,
               "created_at": __import__("datetime").datetime.now(
                   __import__("datetime").timezone.utc).isoformat()}
    update_user_context(phone, "pending_drive", pending)
    return pending


def test_gate_non_confirmation_never_applies():
    _stage()
    with patch("app.handlers.pending_drive_handler.send_whatsapp_message"):
        consumed = handle_pending_drive(_ib("what's the weather"), USER)
    assert consumed is False
    assert get_user_context(USER["phone_number"]).get("pending_drive") is None


def test_gate_abort_acknowledges_no_change():
    _stage()
    with patch("app.handlers.pending_drive_handler.send_whatsapp_message") as sw:
        consumed = handle_pending_drive(_ib("no, cancel"), USER)
    assert consumed is True
    assert "didn't change" in sw.call_args[0][1].lower()


def test_gate_confirmation_dispatches_apply():
    _stage()
    with patch("app.agents.drive_agent.DriveAgent.run_skill") as rs, \
         patch("app.handlers.pending_drive_handler.send_whatsapp_message"), \
         patch("app.handlers.pending_drive_handler.format_response", return_value="Done"):
        rs.return_value = AgentResult(agent_name="DriveAgent", success=True,
                                      data={"type": "drive_modify_applied",
                                            "file_name": "Ventas"})
        consumed = handle_pending_drive(_ib("yes"), USER)
    assert consumed is True
    assert rs.call_args[0][0] == "apply_modification"


# ── apply: revision guard ───────────────────────────────────────────────────

def test_apply_writes_when_revision_matches():
    pend = _stage()
    with patch("app.agents.drive_agent.skills.apply_modification.get_drive_refresh_token",
               return_value="rt"), \
         patch("app.services.google_drive.get_file_meta",
               return_value={"headRevisionId": "rev1"}), \
         patch("app.services.google_drive.update_sheet_cell") as wrote:
        r = DriveAgent().run_skill("apply_modification",
                                   SkillContext(user=USER, inbound_text="yes",
                                                payload=pend))
    assert r.success and r.data["type"] == "drive_modify_applied"
    wrote.assert_called_once_with("rt", "S1", "Hoja1", "B2", "OK")


def test_apply_blocks_write_on_revision_drift():
    pend = _stage()
    with patch("app.agents.drive_agent.skills.apply_modification.get_drive_refresh_token",
               return_value="rt"), \
         patch("app.services.google_drive.get_file_meta",
               return_value={"headRevisionId": "DIFFERENT"}), \
         patch("app.services.google_drive.update_sheet_cell") as wrote:
        r = DriveAgent().run_skill("apply_modification",
                                   SkillContext(user=USER, inbound_text="yes",
                                                payload=pend))
    assert r.success and r.data["type"] == "drive_modify_revision_conflict"
    wrote.assert_not_called()


# ── routing ─────────────────────────────────────────────────────────────────

def test_router_routes_drive_noun_plus_action():
    d = route(_pm("lee mi documento Ventas en drive"))
    assert d.agent.__class__.__name__ == "DriveAgent"


def test_router_routes_drive_intent_from_parser():
    d = route(_pm("change it", drive_intent="modify", drive_file_ref="Ventas",
                   drive_edit={"op": "append_text", "text": "x"}))
    assert d.agent.__class__.__name__ == "DriveAgent"


@pytest.mark.parametrize("text", [
    "hola", "cuánto gasté esta semana", "qué clima hace hoy",
    "tengo reunión mañana", "gracias",
])
def test_router_non_drive_unaffected(text):
    d = route(_pm(text))
    assert d.agent is not None
    assert d.agent.__class__.__name__ != "DriveAgent"


# ── uploaded Office files: read-only support ────────────────────────────────

def _xlsx_bytes(rows):
    import io
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_bytes(paragraphs):
    import io
    from docx import Document
    doc = Document()
    for p in paragraphs:
        doc.add_paragraph(p)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_xlsx_parsed_to_csv():
    from app.services.google_drive import _xlsx_to_csv
    out = _xlsx_to_csv(_xlsx_bytes([["cliente", "estado"], ["pepito", "ok"]]))
    assert "cliente,estado" in out and "pepito,ok" in out


def test_docx_parsed_to_text():
    from app.services.google_drive import _docx_to_text
    out = _docx_to_text(_docx_bytes(["Informe Abril", "Total: 2000000"]))
    assert "Informe Abril" in out and "Total: 2000000" in out


def test_get_content_routes_xlsx_and_docx():
    from app.services import google_drive
    with patch("app.services.google_drive.get_drive_service_for_user",
               return_value=MagicMock()), \
         patch("app.services.google_drive._download_bytes",
               return_value=_xlsx_bytes([["a", "b"], ["1", "2"]])):
        out = google_drive.get_content("rt", "F1", google_drive.XLSX)
    assert "a,b" in out and "1,2" in out


def test_office_files_have_no_write_path():
    """Read-only guarantee: a modify request against an uploaded .xlsx is
    refused (edit_unsupported_for_type) and NOTHING is written."""
    pm = _pm("set estado OK for cliente pepito in Prueba",
             drive_intent="modify", drive_file_ref="Prueba",
             drive_edit={"op": "set_cell", "locator_column": "cliente",
                         "locator_value": "pepito", "target_column": "estado",
                         "new_value": "OK"})
    XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    update_user_context(USER["phone_number"], "pending_drive", None)
    with patch("app.agents.drive_agent._shared.drive_client.decrypt", return_value="rt"), \
         patch("app.services.google_drive.search_files",
               return_value=[{"id": "X1", "name": "Prueba", "mimeType": XLSX}]), \
         patch("app.services.google_drive.update_sheet_cell") as wrote:
        res = DriveAgent().execute(pm, USER)
    assert res.error_message == "edit_unsupported_for_type"
    wrote.assert_not_called()
    assert get_user_context(USER["phone_number"]).get("pending_drive") is None


# ── connect-link side effect ────────────────────────────────────────────────

def test_not_connected_sends_link_and_stays_silent():
    pm = _pm("read Ventas", drive_intent="read", drive_file_ref="Ventas")
    with patch("app.agents.drive_agent.agent.send_connect_link") as link:
        res = DriveAgent().execute(pm, {"phone_number": "+57", "language": "en"})
    assert res.success and res.data["type"] == "drive_connect_link_sent"
    assert link.called
    assert format_response(res, {"language": "en"}) == ""
